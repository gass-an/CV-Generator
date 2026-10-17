import uuid
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

import pytest
from app.clients.avp_client import (
    AvpClientError,
    AvpData,
    AvpInvalidResponseError,
    AvpNotFoundError,
)
from app.clients.llm_client import (
    LlmClientError,
    LlmInvalidResponseError,
    LlmTimeoutError,
)
from app.models.document_job import DocumentType
from app.processors.cover_letter_document_processor import (
    CoverLetterDocumentProcessor,
    include_offer_reference_in_subject,
)
from app.processors.document_processor import DocumentJobData, DocumentProcessingError
from app.renderers import DocxRenderer
from docx import Document

AVP = AvpData(
    reference="3134-26-1382/SR",
    content="# Missions\n- Conduite de projet",
    source_url="https://opt.example/index.md",
    is_archived=False,
)
REFERENCE_TIME = datetime(2026, 10, 17, 1, 30, tzinfo=UTC)
RESUME = {"skills": [{"name": "Python"}]}


class FakeAvpClient:
    def __init__(self, result: AvpData | Exception) -> None:
        self.result = result
        self.requested_number: str | None = None

    async def get_avp(self, avp_number: str) -> AvpData:
        self.requested_number = avp_number
        if isinstance(self.result, Exception):
            raise self.result
        return self.result


class RecordingPromptBuilder:
    def __init__(self) -> None:
        self.resume_data: dict[str, Any] | None = None
        self.avp: AvpData | None = None

    def build(
        self, *, resume_data: dict[str, Any], avp: AvpData
    ) -> list[dict[str, str]]:
        self.resume_data = resume_data
        self.avp = avp
        return [{"role": "user", "content": "sources séparées"}]


class FakeLlmClient:
    def __init__(self, result: str | Exception) -> None:
        self.result = result

    async def generate(self, messages: list[dict[str, str]]) -> str:
        if isinstance(self.result, Exception):
            raise self.result
        return self.result


def make_job() -> DocumentJobData:
    return DocumentJobData(
        id=uuid.uuid4(),
        document_type=DocumentType.COVER_LETTER,
        avp_number=AVP.reference,
        resume_data=RESUME,
        started_at=REFERENCE_TIME,
    )


def make_processor(
    *,
    avp_result: AvpData | Exception = AVP,
    llm_result: str | Exception = "= Lettre de motivation\n\nContenu factuel.",
) -> tuple[CoverLetterDocumentProcessor, FakeAvpClient, RecordingPromptBuilder]:
    avp_client = FakeAvpClient(avp_result)
    builder = RecordingPromptBuilder()
    processor = CoverLetterDocumentProcessor(
        avp_client=avp_client,  # type: ignore[arg-type]
        prompt_builder=builder,  # type: ignore[arg-type]
        llm_client=FakeLlmClient(llm_result),  # type: ignore[arg-type]
    )
    return processor, avp_client, builder


@pytest.mark.asyncio
async def test_process_returns_asciidoc_and_keeps_sources_separate() -> None:
    processor, avp_client, builder = make_processor()

    result = await processor.process(make_job())

    assert result.format == "asciidoc"
    assert result.content.startswith("= Lettre de motivation")
    assert avp_client.requested_number == AVP.reference
    assert builder.resume_data is RESUME
    assert builder.avp is AVP
    assert "Conduite de projet" not in str(builder.resume_data)


@pytest.mark.asyncio
async def test_process_adds_offer_reference_to_subject() -> None:
    processor, _, _ = make_processor(
        llm_result=(
            "= Lettre de motivation\n\n"
            "Objet : Candidature au poste d'administrateur\n\nMadame, Monsieur,"
        )
    )

    result = await processor.process(make_job())

    assert (
        "Objet : Candidature à l'offre n° 3134-26-1382/SR – administrateur"
        in result.content
    )
    rendered_docx = Document(BytesIO(DocxRenderer().render(result.content)))
    assert any(
        paragraph.text
        == "Objet : Candidature à l'offre n° 3134-26-1382/SR – administrateur"
        for paragraph in rendered_docx.paragraphs
    )


def test_subject_is_unchanged_without_offer_reference() -> None:
    content = "= Lettre de motivation\n\nObjet : Candidature spontanée"

    assert include_offer_reference_in_subject(content, "  ") == content


@pytest.mark.parametrize(
    ("empty_subject", "expected_subject"),
    [
        ("Objet :", "Objet : Candidature à l'offre n° 3134-26-1382/SR"),
        (
            "*Objet :*",
            "*Objet :* Candidature à l'offre n° 3134-26-1382/SR",
        ),
        (
            "**Objet :**",
            "**Objet :** Candidature à l'offre n° 3134-26-1382/SR",
        ),
    ],
)
def test_empty_subject_is_completed_without_absorbing_salutation(
    empty_subject: str, expected_subject: str
) -> None:
    salutation = "Madame, Monsieur,"
    content = f"= Lettre de motivation\n\n{empty_subject}\n{salutation}\n\nTexte."

    once = include_offer_reference_in_subject(content, "3134-26-1382/SR")
    twice = include_offer_reference_in_subject(once, "3134-26-1382/SR")

    assert once == (
        f"= Lettre de motivation\n\n{expected_subject}\n{salutation}\n\nTexte."
    )
    assert twice == once
    assert once.count("Objet :") == 1
    assert once.count("3134-26-1382/SR") == 1
    assert once.count(salutation) == 1


@pytest.mark.parametrize(
    ("subject_line", "expected_line"),
    [
        (
            "Objet : Candidature au poste de développeur",
            "Objet : Candidature à l'offre n° 3134-26-1382/SR – développeur",
        ),
        (
            "*Objet :* Candidature au poste de développeur",
            "*Objet :* Candidature à l'offre n° 3134-26-1382/SR – développeur",
        ),
        (
            "*Objet : Candidature au poste de développeur*",
            "*Objet : Candidature à l'offre n° 3134-26-1382/SR – développeur*",
        ),
        (
            "**Objet : Candidature au poste de développeur**",
            "**Objet : Candidature à l'offre n° 3134-26-1382/SR – développeur**",
        ),
        (
            "Objet : Demande de mobilité – développeur",
            "Objet : Demande de mobilité – développeur – offre n° 3134-26-1382/SR",
        ),
        (
            "Objet : Candidature au poste d’administrateur d’application",
            "Objet : Candidature à l'offre n° 3134-26-1382/SR"
            " – administrateur d’application",
        ),
    ],
)
def test_subject_variants_are_enriched_once_and_keep_style_and_title(
    subject_line: str, expected_line: str
) -> None:
    content = f"= Lettre de motivation\n\n{subject_line}\n\nMadame, Monsieur,"

    once = include_offer_reference_in_subject(content, "3134-26-1382/SR")
    twice = include_offer_reference_in_subject(once, "3134-26-1382/SR")

    assert expected_line in once
    assert once == twice
    assert once.count("Objet :") == 1
    assert once.count("3134-26-1382/SR") == 1


def test_real_fully_bold_subject_is_updated_without_duplicate() -> None:
    content = (
        "= Lettre de motivation\n\n"
        "Quoc-Kim BUI\n"
        "qkimbui@gmail.com\n\n"
        "*Objet : Candidature au poste d’Administrateur des pratiques métiers*\n\n"
        "Madame, Monsieur,\n\n"
        "Actuellement en Master MIAGE, je souhaite postuler..."
    )

    once = include_offer_reference_in_subject(content, "3134-26-1382/SR")
    twice = include_offer_reference_in_subject(once, "3134-26-1382/SR")

    assert (
        "*Objet : Candidature à l'offre n° 3134-26-1382/SR"
        " – Administrateur des pratiques métiers*"
    ) in once
    assert once == twice
    assert once.count("Objet :") == 1
    assert once.count("3134-26-1382/SR") == 1
    assert "Administrateur des pratiques métiers" in once
    assert once.index("Objet :") < once.index("Madame, Monsieur,")

    rendered_docx = Document(BytesIO(DocxRenderer().render(once)))
    subject_paragraphs = [
        paragraph
        for paragraph in rendered_docx.paragraphs
        if paragraph.text.startswith("Objet :")
    ]
    assert len(subject_paragraphs) == 1
    assert subject_paragraphs[0].text == (
        "Objet : Candidature à l'offre n° 3134-26-1382/SR"
        " – Administrateur des pratiques métiers"
    )
    assert subject_paragraphs[0].runs
    assert all(run.bold for run in subject_paragraphs[0].runs)


def test_subject_does_not_duplicate_existing_offer_reference() -> None:
    content = (
        "= Lettre de motivation\n\n"
        "*Objet :* Candidature à l'offre n° 3134-26-1382/SR – Administrateur"
    )

    result = include_offer_reference_in_subject(content, "3134-26-1382/SR")

    assert result == content
    assert result.count("3134-26-1382/SR") == 1


def test_missing_subject_is_inserted_without_inventing_title() -> None:
    content = (
        "= Lettre de motivation\n\n"
        "Camille Exemple\n"
        "camille@example.nc\n\n"
        "Madame, Monsieur,\n\n"
        "Je vous adresse ma candidature."
    )

    result = include_offer_reference_in_subject(content, "3134-26-1382/SR")

    assert result == (
        "= Lettre de motivation\n\n"
        "Camille Exemple\n"
        "camille@example.nc\n\n"
        "Objet : Candidature à l'offre n° 3134-26-1382/SR\n\n"
        "Madame, Monsieur,\n\n"
        "Je vous adresse ma candidature."
    )
    assert "administrateur" not in result.casefold()
    assert include_offer_reference_in_subject(result, "3134-26-1382/SR") == result
    assert result.count("Objet :") == 1
    assert result.count("3134-26-1382/SR") == 1


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("error", "code"),
    [
        (AvpNotFoundError(), "avp_not_found"),
        (AvpInvalidResponseError(), "avp_invalid_response"),
        (AvpClientError(), "avp_source_error"),
    ],
)
async def test_process_maps_avp_errors(error: Exception, code: str) -> None:
    processor, _, _ = make_processor(avp_result=error)
    with pytest.raises(DocumentProcessingError) as caught:
        await processor.process(make_job())
    assert caught.value.code == code


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("error", "code"),
    [
        (LlmTimeoutError(), "llm_timeout"),
        (LlmClientError(), "llm_api_error"),
        (LlmInvalidResponseError(), "llm_invalid_response"),
    ],
)
async def test_process_maps_llm_errors(error: Exception, code: str) -> None:
    processor, _, _ = make_processor(llm_result=error)
    with pytest.raises(DocumentProcessingError) as caught:
        await processor.process(make_job())
    assert caught.value.code == code


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "content",
    ["", "   ", "```asciidoc\n= Lettre de motivation\n```"],
)
async def test_process_rejects_invalid_generated_document(content: str) -> None:
    processor, _, _ = make_processor(llm_result=content)
    with pytest.raises(DocumentProcessingError) as caught:
        await processor.process(make_job())
    assert caught.value.code == "invalid_generated_document"
