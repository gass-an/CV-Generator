from io import BytesIO

import pytest
from app.renderers import DocxRenderer, DocxRenderingError
from docx import Document

ASCIIDOC_CV = """= Camille Exemple

Email : camille@example.nc
Téléphone : +687 75 00 00

== Profil

Professionnelle organisée et attentive à la qualité.

== Compétences

* Rédaction de procédures : Documentation, Formalisation
* Analyse de processus : Amélioration continue, Organisation

== Expériences professionnelles

=== Assistante qualité — Entreprise Exemple

2023 - 2025

* Rédaction et mise à jour de procédures internes
* Analyse de processus existants

== Formation

=== Licence en Gestion — Université Exemple

2018 - 2021
"""


def test_render_creates_readable_docx_with_expected_content_and_styles() -> None:
    rendered = DocxRenderer().render(ASCIIDOC_CV)

    document = Document(BytesIO(rendered))
    paragraphs = {
        paragraph.text: paragraph.style.name for paragraph in document.paragraphs
    }

    assert paragraphs["Camille Exemple"] == "Title"
    assert paragraphs["Profil"] == "Heading 1"
    assert paragraphs["Compétences"] == "Heading 1"
    assert paragraphs["Expériences professionnelles"] == "Heading 1"
    assert paragraphs["Assistante qualité — Entreprise Exemple"] == "Heading 2"
    assert paragraphs["Formation"] == "Heading 1"
    assert paragraphs["Licence en Gestion — Université Exemple"] == "Heading 2"
    assert (
        paragraphs["Rédaction de procédures : Documentation, Formalisation"]
        == "List Bullet"
    )
    assert paragraphs["Analyse de processus existants"] == "List Bullet"


def test_render_normalizes_legacy_markdown_and_applies_inline_styles() -> None:
    legacy_content = """# Guide de préparation

## Points forts

Un paragraphe avec **un atout** et _un conseil_ distincts.

- **Compétence :** service en salle
  - Preuve documentée
- Autre élément
"""

    document = Document(BytesIO(DocxRenderer().render(legacy_content)))
    paragraphs = document.paragraphs

    assert paragraphs[0].text == "Guide de préparation"
    assert paragraphs[0].style.name == "Title"
    assert paragraphs[1].text == "Points forts"
    assert paragraphs[1].style.name == "Heading 1"
    assert paragraphs[2].style.name == "Normal"
    assert any(run.text == "un atout" and run.bold for run in paragraphs[2].runs)
    assert any(run.text == "un conseil" and run.italic for run in paragraphs[2].runs)
    assert all("**" not in paragraph.text for paragraph in paragraphs)

    competence = next(p for p in paragraphs if p.text.startswith("Compétence"))
    nested = next(p for p in paragraphs if p.text == "Preuve documentée")
    other = next(p for p in paragraphs if p.text == "Autre élément")
    assert competence.style.name == "List Bullet"
    assert any(run.text == "Compétence :" and run.bold for run in competence.runs)
    assert nested.style.name == "List Bullet 2"
    assert other.style.name == "List Bullet"
    assert nested.paragraph_format.left_indent > competence.paragraph_format.left_indent


def test_render_supports_asciidoc_numbered_lists_and_separate_paragraphs() -> None:
    content = """= Exemple

Premier paragraphe.

Deuxième paragraphe.

. Première étape
.. Sous-étape
"""
    document = Document(BytesIO(DocxRenderer().render(content)))
    by_text = {paragraph.text: paragraph for paragraph in document.paragraphs}

    assert by_text["Premier paragraphe."].style.name == "Normal"
    assert by_text["Deuxième paragraphe."].style.name == "Normal"
    assert by_text["Première étape"].style.name == "List Number"
    assert by_text["Sous-étape"].style.name == "List Number 2"


def test_render_normalizes_legacy_markdown_numbered_lists() -> None:
    content = "1. Première étape\n2. Deuxième étape\n  1. Sous-étape"
    document = Document(BytesIO(DocxRenderer().render(content)))
    by_text = {paragraph.text: paragraph for paragraph in document.paragraphs}

    assert by_text["Première étape"].style.name == "List Number"
    assert by_text["Deuxième étape"].style.name == "List Number"
    assert by_text["Sous-étape"].style.name == "List Number 2"
    assert (
        by_text["Sous-étape"].paragraph_format.left_indent
        > by_text["Deuxième étape"].paragraph_format.left_indent
    )


def test_render_preserves_deeper_heading_levels() -> None:
    content = "= Titre\n\n== Niveau un\n\n=== Niveau deux\n\n==== Niveau trois"
    document = Document(BytesIO(DocxRenderer().render(content)))

    assert [paragraph.style.name for paragraph in document.paragraphs] == [
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
    ]


@pytest.mark.parametrize("asciidoc", ["", "   ", "\n\t\n"])
def test_render_rejects_empty_content(asciidoc: str) -> None:
    with pytest.raises(DocxRenderingError, match="ne doit pas être vide"):
        DocxRenderer().render(asciidoc)
