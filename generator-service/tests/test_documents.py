import uuid
from datetime import UTC, datetime, timedelta
from io import BytesIO

import httpx
import pytest
from app.main import app
from app.models.document_job import DocumentJobStatus, DocumentType
from app.services.pdf_conversion_service import (
    PdfConversionError,
    PdfConversionTimeoutError,
    get_pdf_converter,
)
from docx import Document

from tests.conftest import FakeDocumentJobService

ASCIIDOC_CV = """= Camille Exemple

Email : camille@example.nc

== Profil

Professionnelle organisée.

== Compétences

* Rédaction de procédures

== Expériences professionnelles

=== Assistante qualité — Entreprise Exemple

2023 - 2025

== Formation

=== Licence en Gestion — Université Exemple
"""

PDF_CONTENT = b"%PDF-1.4\n1 0 obj\nendobj\n%%EOF"


class FakePdfConverter:
    def __init__(self, result: bytes | Exception = PDF_CONTENT) -> None:
        self.result = result
        self.docx_content: bytes | None = None

    async def convert(self, docx_content: bytes) -> bytes:
        self.docx_content = docx_content
        if isinstance(self.result, Exception):
            raise self.result
        return self.result


@pytest.mark.asyncio
async def test_create_cv_document_returns_pending_job(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    response = await api_client.post(
        "/api/v1/documents/cv",
        json={
            "avp_number": "1234-26-001",
            "resume": {
                "basics": {},
                "work": [],
                "education": [],
                "skills": [],
            },
        },
    )

    assert response.status_code == 202
    payload = response.json()
    assert payload["status"] == "pending"
    created_job = document_service.jobs[uuid.UUID(payload["id"])]
    assert created_job.status is DocumentJobStatus.PENDING
    assert created_job.avp_number == "1234-26-001"


@pytest.mark.asyncio
async def test_create_cover_letter_returns_pending_job_without_exposing_resume(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    resume = {"basics": {"name": "Donnée privée"}, "work": []}

    response = await api_client.post(
        "/api/v1/documents/cover-letter",
        json={"avp_number": " 1234-26-001 ", "resume": resume},
    )

    assert response.status_code == 202
    assert response.json()["status"] == "pending"
    assert "resume" not in response.json()
    assert "Donnée privée" not in response.text
    job = document_service.jobs[uuid.UUID(response.json()["id"])]
    assert job.document_type is DocumentType.COVER_LETTER
    assert job.avp_number == "1234-26-001"
    assert job.resume_data == resume


@pytest.mark.asyncio
async def test_create_interview_prep_returns_owned_pending_job(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    resume = {"basics": {"name": "Malia Exemple"}, "work": []}
    response = await api_client.post(
        "/api/v1/documents/interview-prep",
        json={"avp_number": " REST-2026-014 ", "resume": resume},
    )

    assert response.status_code == 202
    job = document_service.jobs[uuid.UUID(response.json()["id"])]
    assert job.document_type is DocumentType.INTERVIEW_PREP
    assert job.avp_number == "REST-2026-014"
    assert job.resume_data == resume


@pytest.mark.asyncio
async def test_document_history_returns_all_statuses_in_stable_descending_order(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    base_time = datetime(2026, 10, 11, 8, tzinfo=UTC)
    jobs = [
        document_service.add_job(
            job_status,
            document_type=document_type,
            resume_data={
                "basics": {
                    "name": " Camille Exemple ",
                    "email": "prive@example.nc",
                }
            },
            created_at=base_time + timedelta(minutes=index),
        )
        for index, (job_status, document_type) in enumerate(
            [
                (DocumentJobStatus.PENDING, DocumentType.CV),
                (DocumentJobStatus.PROCESSING, DocumentType.COVER_LETTER),
                (DocumentJobStatus.COMPLETED, DocumentType.INTERVIEW_PREP),
                (DocumentJobStatus.FAILED, DocumentType.CV),
            ]
        )
    ]

    response = await api_client.get("/api/v1/documents")

    assert response.status_code == 200
    payload = response.json()
    assert payload["total"] == 4
    assert payload["limit"] == 20
    assert payload["offset"] == 0
    assert [item["id"] for item in payload["items"]] == [
        str(job.id) for job in reversed(jobs)
    ]
    assert {item["status"] for item in payload["items"]} == {
        "pending",
        "processing",
        "completed",
        "failed",
    }
    assert all(item["candidate_name"] == "Camille Exemple" for item in payload["items"])
    assert "email" not in response.text
    assert "resume" not in response.text
    assert "result_content" not in response.text


@pytest.mark.asyncio
async def test_document_history_filters_paginates_and_reports_unpaged_total(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    base_time = datetime(2026, 10, 11, 8, tzinfo=UTC)
    matching_jobs = [
        document_service.add_job(
            DocumentJobStatus.COMPLETED,
            document_type=DocumentType.CV,
            created_at=base_time + timedelta(minutes=index),
            resume_data={"basics": {"name": 42}},
        )
        for index in range(3)
    ]
    document_service.add_job(
        DocumentJobStatus.FAILED,
        document_type=DocumentType.CV,
        created_at=base_time + timedelta(minutes=4),
    )
    document_service.add_job(
        DocumentJobStatus.COMPLETED,
        document_type=DocumentType.COVER_LETTER,
        created_at=base_time + timedelta(minutes=5),
    )

    response = await api_client.get(
        "/api/v1/documents",
        params={
            "document_type": "cv",
            "status": "completed",
            "limit": 1,
            "offset": 1,
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["total"] == 3
    assert payload["limit"] == 1
    assert payload["offset"] == 1
    assert payload["items"][0]["id"] == str(matching_jobs[1].id)
    assert payload["items"][0]["candidate_name"] is None


@pytest.mark.asyncio
async def test_document_history_filters_exactly_by_avp(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    matching = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        avp_number="3134-26-1382/SR",
    )
    document_service.add_job(
        DocumentJobStatus.COMPLETED,
        avp_number="3134-26-1382/SR-AUTRE",
    )

    response = await api_client.get(
        "/api/v1/documents",
        params={"avp_number": "3134-26-1382/SR"},
    )

    assert response.status_code == 200
    assert response.json()["total"] == 1
    assert response.json()["items"][0]["id"] == str(matching.id)


@pytest.mark.asyncio
async def test_document_history_matches_avp_stored_with_peripheral_spaces(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    matching = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        avp_number="  3134-26-1382/SR  ",
    )

    suggestions = await api_client.get("/api/v1/documents/avp-numbers")
    history = await api_client.get(
        "/api/v1/documents",
        params={"avp_number": "3134-26-1382/SR"},
    )

    assert suggestions.json()["items"] == ["3134-26-1382/SR"]
    assert history.status_code == 200
    assert history.json()["total"] == 1
    assert history.json()["items"][0]["id"] == str(matching.id)


@pytest.mark.asyncio
async def test_avp_numbers_are_distinct_sorted_and_filterable_by_prefix(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    for avp_number in ["b", "a", "az", "a", "BN", "cv"]:
        document_service.add_job(
            DocumentJobStatus.COMPLETED,
            avp_number=avp_number,
        )

    all_response = await api_client.get("/api/v1/documents/avp-numbers")
    prefix_response = await api_client.get(
        "/api/v1/documents/avp-numbers",
        params={"prefix": " A "},
    )

    assert all_response.status_code == 200
    assert all_response.json() == {"items": ["a", "az", "b", "BN", "cv"]}
    assert prefix_response.status_code == 200
    assert prefix_response.json() == {"items": ["a", "az"]}


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "params",
    [
        {"limit": 0},
        {"limit": 101},
        {"offset": -1},
        {"document_type": "unknown"},
        {"status": "unknown"},
    ],
)
async def test_document_history_rejects_invalid_parameters(
    api_client: httpx.AsyncClient,
    params: dict[str, int | str],
) -> None:
    response = await api_client.get("/api/v1/documents", params=params)
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_document_history_is_empty_for_client_without_documents(
    api_client: httpx.AsyncClient,
) -> None:
    response = await api_client.get("/api/v1/documents")
    assert response.status_code == 200
    assert response.json() == {"items": [], "total": 0, "limit": 20, "offset": 0}


@pytest.mark.asyncio
async def test_get_status_does_not_expose_resume(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(DocumentJobStatus.PENDING)

    response = await api_client.get(f"/api/v1/documents/{job.id}/status")

    assert response.status_code == 200
    assert response.json()["status"] == "pending"
    assert "resume" not in response.json()
    assert "resume_data" not in response.json()


@pytest.mark.asyncio
async def test_unknown_job_returns_not_found(api_client: httpx.AsyncClient) -> None:
    response = await api_client.get(f"/api/v1/documents/{uuid.uuid4()}/status")

    assert response.status_code == 404
    assert response.json()["detail"]["code"] == "document_job_not_found"
    assert response.json()["detail"]["message"] == "Le document est introuvable"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "job_status",
    [DocumentJobStatus.PENDING, DocumentJobStatus.PROCESSING],
)
async def test_unfinished_job_result_returns_conflict(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
    job_status: DocumentJobStatus,
) -> None:
    job = document_service.add_job(job_status)

    response = await api_client.get(f"/api/v1/documents/{job.id}")

    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "document_not_ready"
    assert response.json()["detail"]["message"] == (
        "La génération du document n'est pas terminée"
    )


@pytest.mark.asyncio
async def test_completed_job_returns_asciidoc(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content="= Exemple",
        result_format="asciidoc",
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}")

    assert response.status_code == 200
    assert response.json() == {
        "id": str(job.id),
        "type": "cv",
        "format": "asciidoc",
        "content": "= Exemple",
    }


@pytest.mark.asyncio
async def test_completed_cover_letter_uses_existing_result_contract(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        document_type=DocumentType.COVER_LETTER,
        result_content="= Lettre de motivation",
        result_format="asciidoc",
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}")

    assert response.status_code == 200
    assert response.json()["type"] == "cover_letter"


@pytest.mark.asyncio
async def test_completed_interview_prep_uses_existing_result_contract(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        document_type=DocumentType.INTERVIEW_PREP,
        result_content="= Guide de préparation à l'entretien",
        result_format="asciidoc",
    )
    response = await api_client.get(f"/api/v1/documents/{job.id}")
    assert response.status_code == 200
    assert response.json()["type"] == "interview_prep"


@pytest.mark.asyncio
async def test_failed_job_returns_safe_business_error(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    private_error = "postgresql://user:secret@postgres/internal SQL failure"
    job = document_service.add_job(
        DocumentJobStatus.FAILED,
        error_message=private_error,
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}")

    assert response.status_code == 409
    assert response.json()["detail"] == {
        "code": "document_generation_failed",
        "message": "La génération du document a échoué",
    }
    assert private_error not in response.text


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "payload",
    [
        {"resume": {}},
        {"avp_number": "", "resume": {}},
        {"avp_number": "   ", "resume": {}},
        {"avp_number": "1" * 129, "resume": {}},
        {"avp_number": "1234-26-001", "resume": []},
    ],
)
async def test_create_cv_document_rejects_invalid_payload(
    api_client: httpx.AsyncClient,
    payload: dict[str, object],
) -> None:
    response = await api_client.post("/api/v1/documents/cv", json=payload)

    assert response.status_code == 422


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "payload",
    [
        {"resume": {}},
        {"avp_number": "   ", "resume": {}},
        {"avp_number": "1234-26-001", "resume": []},
    ],
)
async def test_create_cover_letter_rejects_invalid_payload(
    api_client: httpx.AsyncClient,
    payload: dict[str, object],
) -> None:
    response = await api_client.post("/api/v1/documents/cover-letter", json=payload)
    assert response.status_code == 422


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "payload",
    [
        {"resume": {}},
        {"avp_number": "   ", "resume": {}},
        {"avp_number": "REST-2026-014", "resume": []},
    ],
)
async def test_create_interview_prep_rejects_invalid_payload(
    api_client: httpx.AsyncClient, payload: dict[str, object]
) -> None:
    response = await api_client.post("/api/v1/documents/interview-prep", json=payload)
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_download_completed_job_returns_readable_docx(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content=ASCIIDOC_CV,
        result_format="asciidoc",
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}/docx")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.headers["content-disposition"] == (
        f'attachment; filename="cv-{job.id}.docx"'
    )
    document = Document(BytesIO(response.content))
    text = [paragraph.text for paragraph in document.paragraphs]
    assert "Camille Exemple" in text
    assert "Profil" in text
    assert "Rédaction de procédures" in text
    assert "Assistante qualité — Entreprise Exemple" in text
    assert b"must not leak" not in response.content


@pytest.mark.asyncio
async def test_download_cover_letter_uses_type_specific_filename(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        document_type=DocumentType.COVER_LETTER,
        result_content="= Lettre de motivation\n\nMadame, Monsieur,\n\nContenu.",
        result_format="asciidoc",
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}/docx")

    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        f'attachment; filename="lettre-motivation-{job.id}.docx"'
    )
    document = Document(BytesIO(response.content))
    assert "Lettre de motivation" in [p.text for p in document.paragraphs]


@pytest.mark.asyncio
async def test_download_interview_prep_uses_type_specific_filename(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        document_type=DocumentType.INTERVIEW_PREP,
        result_content=(
            "= Guide de préparation à l'entretien\n\n== Checklist\n\n* Réviser"
        ),
        result_format="asciidoc",
    )
    response = await api_client.get(f"/api/v1/documents/{job.id}/docx")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        f'attachment; filename="guide-entretien-{job.id}.docx"'
    )
    document = Document(BytesIO(response.content))
    assert "Guide de préparation à l'entretien" in [p.text for p in document.paragraphs]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("document_type", "prefix"),
    [
        (DocumentType.CV, "cv"),
        (DocumentType.COVER_LETTER, "lettre-motivation"),
        (DocumentType.INTERVIEW_PREP, "guide-entretien"),
    ],
)
async def test_adoc_returns_exact_utf8_stored_source_for_all_document_types(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
    document_type: DocumentType,
    prefix: str,
) -> None:
    source = "= Élodie\r\n\r\n== Compétences\r\n\r\n** Liste imbriquée — &\r\n"
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        document_type=document_type,
        result_content=source,
        result_format="asciidoc",
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}/adoc")

    assert response.status_code == 200
    assert response.content == source.encode()
    assert response.headers["content-type"] == "text/asciidoc; charset=utf-8"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="{prefix}-{job.id}.adoc"'
    )
    assert response.headers["cache-control"] == "private, no-store"


@pytest.mark.asyncio
async def test_markdown_converts_final_asciidoc_source(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    source = "= CV d'Élodie\n\n== Profil\n\n* *Fiable*\n** Autonome\n. Étape\n.. Détail"
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content=source,
        result_format="asciidoc",
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}/md")

    assert response.status_code == 200
    assert (
        response.text
        == """# CV d'Élodie

## Profil

- **Fiable**
    - Autonome
1. Étape
    1. Détail"""
    )
    assert response.headers["content-type"] == "text/markdown; charset=utf-8"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="cv-{job.id}.md"'
    )
    assert response.headers["cache-control"] == "private, no-store"


@pytest.mark.asyncio
async def test_markdown_conversion_error_is_controlled(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content=ASCIIDOC_CV,
        result_format="asciidoc",
    )

    def fail_conversion(_renderer: object, _source: str) -> str:
        from app.renderers import MarkdownRenderingError

        raise MarkdownRenderingError("détail interne")

    monkeypatch.setattr(
        "app.api.v1.endpoints.documents.MarkdownRenderer.render", fail_conversion
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}/md")

    assert response.status_code == 500
    assert response.json()["detail"] == {
        "code": "markdown_conversion_failed",
        "message": "La conversion du document en Markdown a échoué",
    }
    assert "détail interne" not in response.text


@pytest.mark.asyncio
async def test_removed_download_route_returns_not_found(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content=ASCIIDOC_CV,
        result_format="asciidoc",
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}/download")

    assert response.status_code == 404


@pytest.mark.asyncio
async def test_pdf_completed_job_returns_inline_pdf_with_private_no_store_cache(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content=ASCIIDOC_CV,
        result_format="asciidoc",
    )
    converter = FakePdfConverter()

    async def override_converter() -> FakePdfConverter:
        return converter

    app.dependency_overrides[get_pdf_converter] = override_converter

    response = await api_client.get(f"/api/v1/documents/{job.id}/pdf")

    assert response.status_code == 200
    assert response.content == PDF_CONTENT
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        f'inline; filename="cv-{job.id}.pdf"'
    )
    assert response.headers["cache-control"] == "private, no-store"
    assert converter.docx_content is not None
    assert converter.docx_content.startswith(b"PK")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("error", "code"),
    [
        (PdfConversionError(), "pdf_conversion_failed"),
        (PdfConversionTimeoutError(), "pdf_conversion_timeout"),
    ],
)
async def test_pdf_conversion_errors_are_controlled(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
    error: Exception,
    code: str,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content=ASCIIDOC_CV,
        result_format="asciidoc",
    )

    async def override_converter() -> FakePdfConverter:
        return FakePdfConverter(error)

    app.dependency_overrides[get_pdf_converter] = override_converter

    response = await api_client.get(f"/api/v1/documents/{job.id}/pdf")

    assert response.status_code == 500
    assert response.json()["detail"]["code"] == code
    assert "LibreOffice" not in response.text


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("job_status", "error_code"),
    [
        (DocumentJobStatus.PENDING, "document_not_ready"),
        (DocumentJobStatus.PROCESSING, "document_not_ready"),
        (DocumentJobStatus.FAILED, "document_generation_failed"),
    ],
)
async def test_pdf_unavailable_job_returns_conflict_without_conversion(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
    job_status: DocumentJobStatus,
    error_code: str,
) -> None:
    job = document_service.add_job(job_status)
    converter = FakePdfConverter()

    async def override_converter() -> FakePdfConverter:
        return converter

    app.dependency_overrides[get_pdf_converter] = override_converter

    response = await api_client.get(f"/api/v1/documents/{job.id}/pdf")

    assert response.status_code == 409
    assert response.json()["detail"]["code"] == error_code
    assert converter.docx_content is None


@pytest.mark.asyncio
async def test_pdf_unknown_job_returns_not_found(
    api_client: httpx.AsyncClient,
) -> None:
    async def override_converter() -> FakePdfConverter:
        return FakePdfConverter()

    app.dependency_overrides[get_pdf_converter] = override_converter
    response = await api_client.get(f"/api/v1/documents/{uuid.uuid4()}/pdf")
    assert response.status_code == 404


@pytest.mark.asyncio
async def test_pdf_unusable_result_returns_conflict(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content="= Document",
        result_format="markdown",
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}/pdf")

    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "document_result_unavailable"


@pytest.mark.asyncio
async def test_download_unknown_job_returns_not_found(
    api_client: httpx.AsyncClient,
) -> None:
    response = await api_client.get(f"/api/v1/documents/{uuid.uuid4()}/docx")

    assert response.status_code == 404
    assert response.json()["detail"]["code"] == "document_job_not_found"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("job_status", "error_code"),
    [
        (DocumentJobStatus.PENDING, "document_not_ready"),
        (DocumentJobStatus.PROCESSING, "document_not_ready"),
        (DocumentJobStatus.FAILED, "document_generation_failed"),
    ],
)
async def test_download_unavailable_job_returns_conflict(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
    job_status: DocumentJobStatus,
    error_code: str,
) -> None:
    job = document_service.add_job(job_status)

    response = await api_client.get(f"/api/v1/documents/{job.id}/docx")

    assert response.status_code == 409
    assert response.json()["detail"]["code"] == error_code
    assert "must not leak" not in response.text


@pytest.mark.asyncio
async def test_download_completed_job_without_usable_result_is_controlled(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content="   ",
        result_format="asciidoc",
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}/docx")

    assert response.status_code == 409
    assert response.json()["detail"] == {
        "code": "document_result_unavailable",
        "message": "Le résultat du document est indisponible",
    }
    assert "must not leak" not in response.text


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("result_content", "result_format"),
    [
        (None, "asciidoc"),
        (ASCIIDOC_CV, "markdown"),
    ],
)
async def test_download_completed_job_with_unavailable_result_returns_conflict(
    api_client: httpx.AsyncClient,
    document_service: FakeDocumentJobService,
    result_content: str | None,
    result_format: str,
) -> None:
    job = document_service.add_job(
        DocumentJobStatus.COMPLETED,
        result_content=result_content,
        result_format=result_format,
    )

    response = await api_client.get(f"/api/v1/documents/{job.id}/docx")

    assert response.status_code == 409
    assert response.json()["detail"] == {
        "code": "document_result_unavailable",
        "message": "Le résultat du document est indisponible",
    }


@pytest.mark.asyncio
async def test_openapi_documents_generation_routes_in_french(
    api_client: httpx.AsyncClient,
) -> None:
    response = await api_client.get("/openapi.json")

    assert response.status_code == 200
    schema = response.json()
    assert schema["info"]["title"] == "Service de génération de documents HackAVP"
    assert (
        "démarre la génération asynchrone d'un document"
        in schema["info"]["description"]
    )
    assert "job" not in schema["info"]["description"].lower()

    expected_paths = {
        "/api/v1/health",
        "/api/v1/documents/cv",
        "/api/v1/documents/cover-letter",
        "/api/v1/documents/interview-prep",
        "/api/v1/documents",
        "/api/v1/documents/avp-numbers",
        "/api/v1/documents/{id}/status",
        "/api/v1/documents/{id}",
        "/api/v1/documents/{id}/docx",
        "/api/v1/documents/{id}/pdf",
        "/api/v1/documents/{id}/adoc",
        "/api/v1/documents/{id}/md",
    }
    assert set(schema["paths"]) == expected_paths
    assert not any("{job_id}" in path for path in schema["paths"])
    assert schema["paths"]["/api/v1/documents/cv"]["post"]["summary"] == (
        "Créer une génération de CV"
    )
    assert (
        schema["paths"]["/api/v1/documents/cover-letter"]["post"]["summary"]
        == "Créer une génération de lettre de motivation"
    )
    assert (
        schema["paths"]["/api/v1/documents/interview-prep"]["post"]["summary"]
        == "Créer un guide de préparation à l'entretien"
    )

    for path in (
        "/api/v1/documents/{id}/status",
        "/api/v1/documents/{id}",
        "/api/v1/documents/{id}/docx",
        "/api/v1/documents/{id}/pdf",
        "/api/v1/documents/{id}/adoc",
        "/api/v1/documents/{id}/md",
    ):
        operation = schema["paths"][path]["get"]
        path_parameter = next(
            parameter
            for parameter in operation["parameters"]
            if parameter["in"] == "path"
        )
        assert path_parameter["name"] == "id"
        assert path_parameter["description"] == "Identifiant unique du document."

    components = schema["components"]["schemas"]
    assert components["DocumentType"]["enum"] == [
        "cv",
        "cover_letter",
        "interview_prep",
    ]
    assert components["DocumentJobStatus"]["enum"] == [
        "pending",
        "processing",
        "completed",
        "failed",
    ]
    request_schema = components["CreateDocumentRequest"]["properties"]
    assert request_schema["avp_number"]["examples"] == ["3134-26-1382/SR"]
    assert request_schema["resume"]["examples"][0]["basics"]["name"] == (
        "Camille Exemple"
    )

    public_operations = [
        operation
        for path in schema["paths"].values()
        for operation in path.values()
        if isinstance(operation, dict) and "responses" in operation
    ]
    public_texts = [schema["info"]["description"]]
    for operation in public_operations:
        public_texts.extend(
            operation.get(field, "") for field in ("summary", "description")
        )
        public_texts.extend(
            response.get("description", "")
            for response in operation["responses"].values()
        )
    public_texts.extend(
        component.get("description", "") for component in components.values()
    )
    assert all("job" not in text.lower() for text in public_texts)
