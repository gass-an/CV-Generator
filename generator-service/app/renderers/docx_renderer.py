import re
from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm, Pt

from app.asciidoc import normalize_asciidoc

UNORDERED_LIST_ITEM = re.compile(r"^(?P<markers>\*+)[ \t]+(?P<text>.+)$")
ORDERED_LIST_ITEM = re.compile(r"^(?P<markers>\.+)[ \t]+(?P<text>.+)$")
ASCIIDOC_HEADING = re.compile(r"^(?P<markers>={1,6})[ \t]+(?P<text>.+)$")
INLINE_STYLE = re.compile(
    r"(?<![\\\w])(?P<marker>[*_])(?=\S)(?P<text>.+?)(?<=\S)(?P=marker)(?!\w)"
)


class DocxRenderingError(Exception):
    """Le contenu AsciiDoc ne peut pas être rendu en document DOCX."""


class DocxRenderer:
    """Convertit le sous-ensemble AsciiDoc produit par le LLM en fichier DOCX."""

    def render(self, asciidoc: str) -> bytes:
        """Produit un paquet DOCX valide en mémoire à partir du contenu AsciiDoc."""
        if not asciidoc.strip():
            raise DocxRenderingError("Le contenu AsciiDoc ne doit pas être vide")

        try:
            document = Document()
            self._configure_document(document)
            self._add_content(document, normalize_asciidoc(asciidoc))

            output = BytesIO()
            document.save(output)
            rendered = output.getvalue()

            # La réouverture détecte un paquet en mémoire incomplet ou invalide.
            Document(BytesIO(rendered))
            return rendered
        except DocxRenderingError:
            raise
        except Exception as error:
            raise DocxRenderingError(
                "Impossible de générer le document DOCX"
            ) from error

    def _configure_document(self, document: Document) -> None:
        section = document.sections[0]
        section.start_type = WD_SECTION.NEW_PAGE
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

        styles = document.styles

        normal = styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.08

        title = styles["Title"]
        title.font.name = "Arial"
        title.font.size = Pt(24)
        title.font.bold = True
        title.paragraph_format.space_after = Pt(12)

        heading_1 = styles["Heading 1"]
        heading_1.font.name = "Arial"
        heading_1.font.size = Pt(15)
        heading_1.font.bold = True
        heading_1.paragraph_format.space_before = Pt(12)
        heading_1.paragraph_format.space_after = Pt(5)

        heading_2 = styles["Heading 2"]
        heading_2.font.name = "Arial"
        heading_2.font.size = Pt(11.5)
        heading_2.font.bold = True
        heading_2.paragraph_format.space_before = Pt(8)
        heading_2.paragraph_format.space_after = Pt(3)

        heading_3 = styles["Heading 3"]
        heading_3.font.name = "Arial"
        heading_3.font.size = Pt(10.5)
        heading_3.font.bold = True
        heading_3.paragraph_format.space_before = Pt(6)
        heading_3.paragraph_format.space_after = Pt(3)

        list_bullet = styles["List Bullet"]
        list_bullet.font.name = "Arial"
        list_bullet.font.size = Pt(10.5)
        list_bullet.paragraph_format.space_after = Pt(3)

        for style_name in (
            "List Bullet 2",
            "List Bullet 3",
            "List Number",
            "List Number 2",
            "List Number 3",
        ):
            style = styles[style_name]
            style.font.name = "Arial"
            style.font.size = Pt(10.5)
            style.paragraph_format.space_after = Pt(3)

    def _add_content(self, document: Document, asciidoc: str) -> None:
        for raw_line in asciidoc.splitlines():
            line = raw_line.rstrip()
            if not line.strip():
                continue

            if heading := ASCIIDOC_HEADING.fullmatch(line):
                self._add_heading(
                    document,
                    heading.group("text"),
                    level=len(heading.group("markers")),
                )
            elif unordered := UNORDERED_LIST_ITEM.fullmatch(line):
                self._add_list_item(
                    document,
                    unordered.group("text"),
                    level=len(unordered.group("markers")),
                    ordered=False,
                )
            elif ordered := ORDERED_LIST_ITEM.fullmatch(line):
                self._add_list_item(
                    document,
                    ordered.group("text"),
                    level=len(ordered.group("markers")),
                    ordered=True,
                )
            else:
                paragraph = document.add_paragraph(style="Normal")
                self._add_inline_content(paragraph, line.strip())

    def _add_heading(self, document: Document, text: str, *, level: int) -> None:
        style_name = "Title" if level == 1 else f"Heading {min(level - 1, 9)}"
        paragraph = document.add_paragraph(style=style_name)
        self._add_inline_content(paragraph, text.strip())
        if level == 1:
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    def _add_list_item(
        self,
        document: Document,
        text: str,
        *,
        level: int,
        ordered: bool,
    ) -> None:
        family = "List Number" if ordered else "List Bullet"
        style_level = min(level, 3)
        style_name = family if style_level == 1 else f"{family} {style_level}"
        paragraph = document.add_paragraph(style=style_name)
        paragraph.paragraph_format.left_indent = Cm(0.63 * level)
        paragraph.paragraph_format.first_line_indent = Cm(-0.32)
        self._add_inline_content(paragraph, text.strip())

    @staticmethod
    def _add_inline_content(paragraph: object, text: str) -> None:
        cursor = 0
        for match in INLINE_STYLE.finditer(text):
            if match.start() > cursor:
                paragraph.add_run(text[cursor : match.start()])
            run = paragraph.add_run(match.group("text"))
            if match.group("marker") == "*":
                run.bold = True
            else:
                run.italic = True
            cursor = match.end()
        if cursor < len(text):
            paragraph.add_run(text[cursor:])
